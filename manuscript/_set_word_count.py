"""Write or refresh the trailing "Word Count." line of a rendered JAMA docx.

Must run AFTER `_post_bold_abstract.py` and AFTER the word count is known, i.e.
last in the post-render chain. The line is not in the .qmd (it is a submission
furniture item, not prose), so every render erases it and it has to be reinjected.

Usage:  python3 _set_word_count.py <docx> <body_words>
"""

import re
import sys

import docx

TEMPLATE = ("Word Count. {n} words of text "
            "(excluding abstract, references, tables, and figure legends).")


def main(path, n):
    n = f"{int(n):,}"
    doc = docx.Document(path)

    # Drop any stale line first, wherever it sits, so repeated runs converge
    # rather than stacking copies.
    for p in list(doc.paragraphs):
        if re.match(r"^\s*Word Count\.\s", p.text or ""):
            p._element.getparent().remove(p._element)

    para = doc.add_paragraph(TEMPLATE.format(n=n))
    run = para.runs[0] if para.runs else para.add_run()
    run.bold = True
    doc.save(path)
    print(f"word count line set: {n}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
